"""Export review workbooks; CSV + manifest remain the authoritative pipeline outputs."""
import argparse
import csv
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export(folder):
    folder = Path(folder)
    workbook = Workbook()
    overview = workbook.active
    overview.title = '先看这里'
    overview.append(['项目', '说明'])
    overview.append(['清洗原则', '保留零值及负电流；未知年份/时间不猜测修复；设施只标原编码；费用未估算。'])
    overview.append(['文件分层', 'DWD是清洗明细，ADS是可用聚合，audit是处理原因，ODS是未改动原文件。'])
    overview.append(['文本存储', '为保持ID和Decimal精度，各页数据按文本导出；计算和大屏以CSV字段说明为准。'])
    overview.append(['日期限制', 'calendar_date为空的会话不能画真实日/月趋势；电池不支持时间曲线。'])
    mapping = {'dwd/stations.csv': '站点明细', 'dwd/sessions.csv': '会话明细', 'dwd/battery.csv': '电池明细',
               'ads/by_station_id.csv': '站点聚合', 'ads/by_facility_type.csv': '设施聚合',
               'ads/by_start_hour.csv': '小时聚合', 'ads/by_weekday.csv': '星期聚合',
               'ads/by_platform.csv': '平台聚合', 'audit/rule_counts.csv': '问题数量'}
    expected = {}
    for filename, title in mapping.items():
        with (folder / filename).open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.reader(stream))
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                cell.data_type = 's'
                cell.number_format = '@'
        expected[title] = [tuple(value or None for value in row) for row in rows]
        overview.append([title, f'{len(rows)-1}行；对应{filename}'])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='245B78')
        for column in range(1, sheet.max_column+1):
            sheet.column_dimensions[get_column_letter(column)].width = 26
        sheet.sheet_view.zoomScale = 85
    overview.column_dimensions['B'].width = 105
    for row in overview:
        row[1].alignment = Alignment(wrap_text=True)
    target = folder/'清洗结果查看.xlsx'
    workbook.save(target)
    check = load_workbook(target, read_only=True)
    for name, rows in expected.items():
        assert list(check[name].values) == rows, name
    check.close()
    print(target)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('folder')
    export(parser.parse_args().folder)
